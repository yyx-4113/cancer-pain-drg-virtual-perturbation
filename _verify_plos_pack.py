#!/usr/bin/env python3
"""Verify the PLOS ONE pack is loss-free and compliant. Exit 0 = pass."""
import os, re, subprocess, sys
from docx import Document
from PIL import Image

HERE = r"D:\2026.10\不同部位癌痛患者对药物敏感性生信分析\cancer_pain_virtual_perturbation"
OUT = os.path.join(HERE, "submission_PLOS_ONE")
MS = os.path.join(HERE, "docs", "manuscript", "manuscript_final.md")
GATE = os.path.join(HERE, "scripts", "98_check_manuscript_numbers.py")

NUM = re.compile(r"\d[\d,]*(?:\.\d+)?")
OK, BAD = [], []

def chk(label, got, exp=True):
    (OK if got == exp else BAD).append(f"{label}: got={got!r} exp={exp!r}")

def docx_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)

def docx_tables(path):
    return len(Document(path).tables)

def tokens(s):
    return set(NUM.findall(s.replace("`", "")))

# ---- exported markdown (what goes into the docx) ----
md = open(MS, encoding="utf-8").read()
md_export = "\n".join(
    ln for ln in md.split("\n")
    if not ln.strip().startswith("![") and ln.strip() != "---"
)

ms_text = docx_text(os.path.join(OUT, "Manuscript.docx"))

# 1. numeric tokens survive both ways
a, b = tokens(md_export), tokens(ms_text)
chk("Manuscript 无丢失数字", sorted(a - b), [])
chk("Manuscript 无凭空新增数字", sorted(b - a), [])

# 2. no Chinese anywhere in submitted files
for fn, t in [("Manuscript.docx", ms_text),
              ("Cover_Letter.docx", docx_text(os.path.join(OUT, "Cover_Letter.docx"))),
              ("S1_Supporting_Information.docx", docx_text(os.path.join(OUT, "S1_Supporting_Information.docx")))]:
    chk(f"[{fn}] 无中文", re.findall(r"[\u4e00-\u9fff]", t), [])

# 3. mandatory strings survive
for needle in [
    "OpenAI GPT-4-class",
    "https://github.com/yyx-4113/cancer-pain-drg-virtual-perturbation",
    "release v1.0.0",
    "3/4,901", "0.0019", "binomial p = 0.0067", "anti-selective",
    "Python 3.13", "1.27", "6.8", "0.052", "44.7",
    "Declaration of generative AI",
]:
    chk(f"Manuscript 含「{needle}」", needle in ms_text)

# 4. table counts (5 in manuscript)
chk("Manuscript 表数 == 5", docx_tables(os.path.join(OUT, "Manuscript.docx")), 5)

# 5. placeholders
for fn, t in [("Manuscript.docx", ms_text),
              ("Cover_Letter.docx", docx_text(os.path.join(OUT, "Cover_Letter.docx")))]:
    left = [m for m in ["[[", "TODO", "TBD", "XXX", "COMPLETE BEFORE SUBMISSION"] if m in t]
    chk(f"[{fn}] 无占位符", left, [])

# 6. figure specs
for n in (1, 2, 3):
    p = os.path.join(OUT, f"Figure{n}.tif")
    with Image.open(p) as im:
        dpi = im.info.get("dpi", (0, 0))[0]
    chk(f"Figure{n}.tif <= 10 MB", os.path.getsize(p) <= 10 * 1024 * 1024)
    chk(f"Figure{n}.tif >= 300 ppi", round(float(dpi)) >= 300)

# 7. gate still 291/0 on the polished source
r = subprocess.run(["C:\\Users\\Administrator\\.workbuddy\\binaries\\python\\envs\\default\\Scripts\\python.exe", GATE],
                   cwd=HERE, capture_output=True, encoding="utf-8", errors="replace")
chk("gate 291/0", "291 checks, 0 failed" in r.stdout, True)

print(f"==== PASS {len(OK)} / FAIL {len(BAD)} ====")
for x in BAD:
    print("  [X]", x)
if BAD:
    # surface the lost/added numeric tokens for diagnosis
    if a - b or b - a:
        print("  lost :", sorted(a - b))
        print("  added:", sorted(b - a))
sys.exit(1 if BAD else 0)
